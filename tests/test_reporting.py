import io

from openpyxl import load_workbook
from pypdf import PdfReader

from credittool_engine.reporting import build_excel_report, build_pdf_report


def _sample_report_data():
    return {
        "project_name": "Test Scorecard",
        "scorecard_type": "application",
        "target_definition": {"summary": {"n_total": 200, "n_good": 168, "n_bad": 32, "n_excluded": 0, "bad_rate_pct": 16.0}},
        "sampling": {"summary": {"train": {"n_total": 140, "bad_rate_pct": 15.7}, "test": {"n_total": 60, "bad_rate_pct": 16.6}}},
        "binning": {
            "variables": {
                "age": {"kind": "numeric", "iv": 0.12, "bins": [{"label": "(-inf, 44.0]", "woe": 0.2, "points": -12}]},
                "segment": {"kind": "categorical", "iv": 0.21, "bins": [{"label": "A", "woe": 0.4, "points": 11}]},
            }
        },
        "modeling": {
            "intercept": -1.68,
            "n_obs": 140,
            "pseudo_r2": 0.05,
            "coefficients": [{"variable": "age", "coefficient": -0.86, "std_error": 0.8, "z_value": -1.07, "p_value": 0.28}],
        },
        "scaling": {
            "base_points": 536,
            "min_score": 490,
            "max_score": 568,
            "variables": [{"variable": "age", "bins": [{"label": "(-inf, 44.0]", "woe": 0.2, "points": -12}]}],
        },
        "validation": {
            "train": {"n_total": 140, "bad_rate_pct": 15.7, "ks": 32.0, "auc": 0.64, "gini": 0.29},
            "test": {"n_total": 60, "bad_rate_pct": 16.6, "ks": 26.0, "auc": 0.41, "gini": -0.17},
            "oot": None,
            "psi_test_vs_train": {"index": 0.29, "buckets": []},
        },
    }


def test_excel_report_has_expected_sheets_and_content():
    report_bytes = build_excel_report(_sample_report_data())
    wb = load_workbook(io.BytesIO(report_bytes))

    assert wb.sheetnames == ["Summary", "Binning & IV", "Model", "Scorecard"]

    binning_ws = wb["Binning & IV"]
    values = [row[0].value for row in binning_ws.iter_rows(min_row=2, max_col=1)]
    assert "age" in values
    assert "segment" in values

    scorecard_ws = wb["Scorecard"]
    assert scorecard_ws["A1"].value == "Base points"
    assert scorecard_ws["B1"].value == 536


def test_pdf_report_builds_and_contains_project_name():
    report_bytes = build_pdf_report(_sample_report_data())
    assert report_bytes[:4] == b"%PDF"

    reader = PdfReader(io.BytesIO(report_bytes))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Test Scorecard" in text
    assert "Validation" in text


def test_pdf_report_handles_missing_optional_sections():
    data = _sample_report_data()
    data["validation"] = None
    report_bytes = build_pdf_report(data)
    assert report_bytes[:4] == b"%PDF"


def _penalised_report_data():
    data = _sample_report_data()
    data["modeling"].update(
        penalty="l1",
        alpha=0.02,
        dropped_variables=["segment"],
        coefficients=[{"variable": "age", "coefficient": -0.5, "std_error": None, "z_value": None, "p_value": None}],
    )
    return data


def test_excel_report_marks_penalised_fit():
    wb = load_workbook(io.BytesIO(build_excel_report(_penalised_report_data())))
    rows = [[c.value for c in row] for row in wb["Model"].iter_rows()]
    assert ["age", -0.5, "n/a", "n/a", "n/a"] in rows
    assert ["Penalty", "L1", "alpha", 0.02] in [r[:4] for r in rows]
    assert ["Dropped by penalty", "segment"] in [r[:2] for r in rows]


def test_pdf_report_marks_penalised_fit():
    reader = PdfReader(io.BytesIO(build_pdf_report(_penalised_report_data())))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "None" not in text
    assert "n/a" in text
    assert "L1 penalty" in text and "segment" in text


def _rich_report_data():
    data = _sample_report_data()
    data["binning"]["variables"]["age"]["bins"] = [
        {"label": "(-inf, 44.0]", "n_total": 80, "n_good": 64, "n_bad": 16, "bad_rate_pct": 20.0, "pct_of_total": 57.14, "woe": -0.27},
        {"label": "(44.0, inf]", "n_total": 60, "n_good": 54, "n_bad": 6, "bad_rate_pct": 10.0, "pct_of_total": 42.86, "woe": 0.52},
    ]
    data["final_model"] = {"source": "logistic", "algorithm": "logistic", "variables": ["age"]}
    data["modeling"]["variables"] = ["age"]
    data["target_definition"]["definition"] = {"target_column": "status", "mode": "values", "bad_values": ["default"]}
    data["sampling"]["config"] = {"test_size": 0.3, "oot_size": 0.2, "date_column": "app_date", "random_state": 42}
    data["validation"]["oot"] = {"n_total": 50, "bad_rate_pct": 20.0, "ks": 20.0, "auc": 0.6, "gini": 0.2}
    data["candidates"] = [
        {"id": "a", "kind": "xgboost", "is_active": True, "variables": ["age", "segment"], "settings": {"n_estimators": 300, "max_depth": 6}, "sample_metrics": {"train": {"gini": 0.9}, "test": {"gini": 0.3}}},
        {"id": "b", "kind": "logistic", "is_active": True, "variables": ["age"], "settings": {"penalty": "none", "alpha": 0}, "sample_metrics": {"train": {"gini": 0.29}}},
    ]
    data["audit"] = [{"created_at": "2026-09-25T06:00:00", "username": "dinesh", "role": "analyst", "action": "project.create", "detail": "Created"}]
    data["variable_selection"] = {
        "config": {"min_iv": 0.02, "max_corr": 0.7, "max_vif": 5.0},
        "selected": ["age"],
        "dropped": [{"variable": "segment", "reason": "correlation 0.81 with age"}],
        "vif": {"age": 1.0},
    }
    data["generated_by"] = "dinesh"
    data["narrative"] = {"purpose_scope": "Approve/decline for personal loans."}
    return data


def test_document_has_the_full_set_of_sections():
    from credittool_engine.reporting import build_document
    from credittool_engine.reporting.document import Heading

    headings = [b.text for b in build_document(_rich_report_data()).blocks if isinstance(b, Heading)]
    for section in ("Executive summary", "Purpose and scope", "Data", "Target definition", "Sampling", "Characteristic analysis",
                    "Model development", "Alternatives considered", "Scorecard", "Validation", "Limitations", "Ongoing monitoring",
                    "Development audit trail", "Sign-off"):
        assert any(section in h for h in headings), section
    assert any(h.startswith("6.1") and "Binning: age" in h for h in headings)  # model variables get binning detail
    assert any("Variable selection" in h for h in headings)


def test_executive_summary_flags_overfitting_and_missing_significance():
    from credittool_engine.reporting import build_document
    from credittool_engine.reporting.document import Bullets

    findings = next(b for b in build_document(_rich_report_data()).blocks if isinstance(b, Bullets)).items
    # test Gini -0.17 vs train 0.29
    assert any(f.startswith("Test Gini is") and "overfitting" in f for f in findings)
    assert any("Score PSI on the test sample is 0.29" in f for f in findings)
    assert any("Not significant at 5%: age" in f for f in findings)


def test_narrative_drafts_use_project_numbers_and_saved_text_wins():
    from credittool_engine.reporting import FIELD_KEYS, draft_narrative, resolve_narrative

    data = _rich_report_data()
    drafts = draft_narrative(data)
    assert set(drafts) == set(FIELD_KEYS)
    assert "status in default" in drafts["target_rationale"] and "16.0%" in drafts["target_rationale"]
    assert "the most recent" not in drafts["limitations"]
    resolved = resolve_narrative(data)
    assert resolved["purpose_scope"] == "Approve/decline for personal loans."
    assert resolved["limitations"] == drafts["limitations"]


def test_word_document_builds_with_toc_headings_and_narrative_shading():
    from docx import Document as WordDocument
    from docx.oxml.ns import qn

    from credittool_engine.reporting import build_docx_report

    word = WordDocument(io.BytesIO(build_docx_report(_rich_report_data())))
    headings = [p.text for p in word.paragraphs if p.style.name.startswith("Heading")]
    assert "1  Executive summary" in headings
    assert 'TOC \\o "1-2"' in word.element.body.xml
    shaded = [p.text for p in word.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:shd")) is not None]
    assert "Approve/decline for personal loans." in shaded
    assert len(word.inline_shapes) == 2  # IV chart + age's bad-rate-by-bin (fixture has no validation curves)


def test_pdf_renders_non_latin_symbols():
    reader = PdfReader(io.BytesIO(build_pdf_report(_rich_report_data())))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "≥" in text and "Executive summary" in text and "Approve/decline for personal loans." in text
