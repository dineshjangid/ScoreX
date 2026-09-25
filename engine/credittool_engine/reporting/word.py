"""The model development document as an editable Word file (python-docx),
rendered from the same structure as the PDF (document.py)."""
from __future__ import annotations

import io

from docx import Document as WordDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from . import document as d
from .document import build_document

INK = RGBColor(0x1F, 0x1D, 0x24)
MUTED = RGBColor(0x6B, 0x63, 0x75)
HEADER_FILL = "F3F2F5"
NARRATIVE_FILL = "F4F7FD"


def _shade(element, fill: str) -> None:
    props = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def _set_cell_margins(table, top: int = 40, bottom: int = 40) -> None:
    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("bottom", bottom)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    # Schema order: tblCellMar comes before tblLook (Word rejects files
    # whose elements are out of order).
    look = tbl_pr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(margins)
    else:
        tbl_pr.append(margins)


def _add_toc(word: WordDocument) -> None:
    paragraph = word.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Table of contents — right-click and choose Update Field if it is empty."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, placeholder, end):
        run._r.append(el)
    # Ask Word to refresh fields (the TOC) when the file is opened.
    settings = word.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    # Schema order: updateFields precedes these settings elements.
    followers = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
                 "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                 "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                 "decimalSymbol", "listSeparator"}
    follower = next((el for el in settings if el.tag.split("}")[-1] in followers), None)
    if follower is not None:
        follower.addprevious(update)
    else:
        settings.append(update)


def _configure_styles(word: WordDocument) -> None:
    normal = word.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 26), ("Heading 1", 15), ("Heading 2", 12)):
        style = word.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = INK
        style.font.bold = name != "Title"
    word.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    word.styles["Heading 2"].paragraph_format.space_before = Pt(10)


def _caption(word: WordDocument, text: str) -> None:
    p = word.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    run.italic = True


def _table(word: WordDocument, block: d.Table) -> None:
    table = word.add_table(rows=1, cols=len(block.headers))
    table.style = word.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_cell_margins(table)
    for i, header in enumerate(block.headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(header))
        run.bold = True
        run.font.size = Pt(8.5)
        _shade(cell._tc, HEADER_FILL)
    # Repeat the header row on each page.
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    tr_pr.append(header_flag)
    for row in block.rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            paragraph = cells[i].paragraphs[0]
            run = paragraph.add_run("" if value is None else str(value))
            run.font.size = Pt(8.5)
            if i in block.numeric:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if block.caption:
        _caption(word, block.caption)
    else:
        word.add_paragraph()


def _key_values(word: WordDocument, block: d.KeyValues) -> None:
    table = word.add_table(rows=0, cols=2)
    _set_cell_margins(table)
    for key, value in block.items:
        cells = table.add_row().cells
        key_run = cells[0].paragraphs[0].add_run(key)
        key_run.font.color.rgb = MUTED
        cells[1].paragraphs[0].add_run(value)
        cells[0].width = Cm(5)
        cells[1].width = Cm(11)
    word.add_paragraph()


def render_docx(doc: d.Document) -> bytes:
    word = WordDocument()
    _configure_styles(word)
    section = word.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    footer = section.footer.paragraphs[0]
    footer.text = f"{doc.title} — {doc.subtitle}"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = MUTED

    word.add_paragraph(doc.title, style="Title")
    subtitle = word.add_paragraph()
    subtitle_run = subtitle.add_run(doc.subtitle)
    subtitle_run.font.size = Pt(15)
    subtitle_run.font.color.rgb = MUTED
    _key_values(word, d.KeyValues(doc.meta))
    word.add_heading("Contents", level=1)
    _add_toc(word)
    word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for block in doc.blocks:
        if isinstance(block, d.Heading):
            word.add_heading(block.text, level=block.level)
        elif isinstance(block, d.Paragraph):
            p = word.add_paragraph(block.text)
            if block.style == "note":
                for run in p.runs:
                    run.italic = True
                    run.font.color.rgb = MUTED
            elif block.style == "narrative":
                # Shaded so reviewers can tell analyst commentary from generated text.
                _shade(p._p, NARRATIVE_FILL)
        elif isinstance(block, d.Bullets):
            for item in block.items:
                word.add_paragraph(item, style="List Bullet")
        elif isinstance(block, d.KeyValues):
            _key_values(word, block)
        elif isinstance(block, d.Table):
            _table(word, block)
        elif isinstance(block, d.Figure):
            word.add_picture(io.BytesIO(block.png), width=Inches(min(block.width_in, 6.3)))
            _caption(word, block.caption)
        elif isinstance(block, d.PageBreak):
            word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def build_docx_report(data: dict) -> bytes:
    return render_docx(build_document(data))
