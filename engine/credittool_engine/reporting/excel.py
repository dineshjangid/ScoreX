"""Multi-sheet Excel workbook documenting a finished scorecard build:
target definition, sampling, binning/IV, selected variables, model
coefficients, the points scorecard, and validation metrics.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FONT = Font(bold=True)


def _or_na(value):
    """Penalised fits have no std error / z / p-value."""
    return "n/a" if value is None else value


def _write_table(ws: Worksheet, headers: list[str], rows: list[list], start_row: int = 1) -> int:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = _HEADER_FONT
    for r, row in enumerate(rows, start=start_row + 1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    return start_row + len(rows) + 2


def build_excel_report(data: dict) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Project", data["project_name"]])
    summary.append(["Scorecard type", data["scorecard_type"]])
    summary.append([])
    target = data.get("target_definition", {}).get("summary", {})
    summary.append(["Target: total", target.get("n_total")])
    summary.append(["Target: good", target.get("n_good")])
    summary.append(["Target: bad", target.get("n_bad")])
    summary.append(["Target: excluded", target.get("n_excluded")])
    summary.append(["Target: bad rate %", target.get("bad_rate_pct")])
    summary.append([])
    sampling = data.get("sampling", {}).get("summary", {})
    for label in ("train", "test", "oot"):
        s = sampling.get(label, {})
        summary.append([f"Sample {label}: n / bad rate %", s.get("n_total"), s.get("bad_rate_pct")])
    validation = data.get("validation") or {}
    if validation:
        summary.append([])
        for label in ("train", "test", "oot"):
            m = validation.get(label)
            if m:
                summary.append([f"{label} KS / AUC / Gini", m.get("ks"), m.get("auc"), m.get("gini")])

    binning_ws = wb.create_sheet("Binning & IV")
    binning_vars = data.get("binning", {}).get("variables", {})
    rows = [[var, b["kind"], b["iv"], len(b["bins"])] for var, b in binning_vars.items()]
    rows.sort(key=lambda r: r[2], reverse=True)
    _write_table(binning_ws, ["Variable", "Kind", "IV", "N bins"], rows)

    final_model = data.get("final_model") or {}
    is_ml_final = final_model.get("source") == "ml_benchmark"

    model_ws = wb.create_sheet("Model")
    if is_ml_final:
        ml_result = data.get("ml_benchmark_result") or {}
        model_ws.append(["Final model", final_model.get("algorithm")])
        model_ws.append(["Train AUC / Gini", ml_result.get("train_auc"), ml_result.get("train_gini")])
        model_ws.append(["Test AUC / Gini", ml_result.get("test_auc"), ml_result.get("test_gini")])
        model_ws.append([])
        importance_rows = [[f["variable"], f["mean_abs_shap"]] for f in ml_result.get("feature_importance", [])]
        _write_table(model_ws, ["Variable", "Mean |SHAP|"], importance_rows, start_row=5)
    else:
        model = data.get("modeling", {})
        model_rows = [["(intercept)", model.get("intercept"), "", "", ""]]
        model_rows += [
            [c["variable"], c["coefficient"], *(_or_na(c.get(k)) for k in ("std_error", "z_value", "p_value"))]
            for c in model.get("coefficients", [])
        ]
        _write_table(model_ws, ["Variable", "Coefficient", "Std error", "z", "p-value"], model_rows)
        if model.get("penalty", "none") != "none":
            model_ws.append([])
            model_ws.append(["Penalty", model["penalty"].upper(), "alpha", model.get("alpha")])
            if model.get("dropped_variables"):
                model_ws.append(["Dropped by penalty", ", ".join(model["dropped_variables"])])

    scorecard_ws = wb.create_sheet("Scorecard")
    if is_ml_final:
        scorecard_ws.append(
            [
                "No points scorecard — final model is a fitted "
                f"{final_model.get('algorithm')} model, deployed via its predicted probability, "
                "not a points table."
            ]
        )
    else:
        scaling = data.get("scaling", {})
        scorecard_ws.append(["Base points", scaling.get("base_points")])
        scorecard_ws.append(["Score range", scaling.get("min_score"), scaling.get("max_score")])
        scorecard_ws.append([])
        points_rows = []
        for v in scaling.get("variables", []):
            for b in v["bins"]:
                points_rows.append([v["variable"], b["label"], b["woe"], b["points"]])
        _write_table(scorecard_ws, ["Variable", "Bin", "WOE", "Points"], points_rows, start_row=4)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
